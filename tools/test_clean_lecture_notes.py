from pathlib import Path
import re
from docx import Document

path = Path('deliverables/Learning Prompt AI - Lecture notes (cleaned).docx')
assert path.exists(), path
doc = Document(path)
text = '\n'.join(p.text for p in doc.paragraphs)
assert not re.search(r'(?<!\d)\d{1,2}:\d{2}(?!\d)', text), 'Time codes remain'
headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
assert len(headings) == 21, headings
assert len({h for h in headings}) == len(headings), 'Duplicate lecture heading remains'
assert any('Pre-trained knowledge' in h for h in headings)
assert any('Data analysis' in h for h in headings)
assert 'Editorial cleanup:' in text
print(f'PASS: cleaned lecture notes contain {len(doc.paragraphs)} paragraphs, {len(headings)} headings, and no time codes.')
