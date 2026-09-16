"""Create a readable, timecode-free copy of the supplied lecture transcript."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt

SOURCE = Path(r"C:/Users/nairi/Downloads/Learning Prompt AI - Lecture notes.docx")
OUTPUT = Path("deliverables/Learning Prompt AI - Lecture notes (cleaned).docx")

TIMECODE = re.compile(r"(?<!\d)\d{1,2}:\d{2}(?!\d)")
SPACE = re.compile(r"\s+")
PART = re.compile(r"^Part\s+(\d+)(?=\s|Lecture|$)", re.I)
LECTURE = re.compile(r"^Lecture\s+(\d+)(?=[A-Z\s]|$)", re.I)

TITLES = {
    (1, 1): "From novice use to capable use",
    (1, 2): "Pre-trained knowledge and its limits",
    (1, 3): "When to use web search",
    (1, 4): "Source quality and verification",
    (1, 5): "Deep research for complex questions",
    (1, 6): "Deep research practice",
    (2, 1): "Brainstorming as a thought partner",
    (2, 2): "Context: what the model can work with",
    (2, 3): "Desktop AI and permission-aware context",
    (2, 4): "Reasoning over complex evidence",
    (2, 5): "Neutral prompting and sycophancy",
    (2, 6): "Writing with AI without generic output",
    (2, 7): "Editing and critique with AI",
    (3, 1): "Multimodal inputs and outputs",
    (3, 2): "Using images as context",
    (3, 3): "Image generation and editing",
    (3, 4): "Building small apps with AI",
    (3, 5): "Data analysis with AI",
    (3, 6): "Practice and final project",
}


def clean(text: str) -> str:
    text = TIMECODE.sub(" ", text)
    text = re.sub(r"\b(Lecture\s+\d+)\s*0{0,2}\s*(?=[A-Z])", r"\1", text)
    text = SPACE.sub(" ", text).strip()
    return text


def is_complete(text: str) -> bool:
    return bool(re.search(r"[.!?…][\"')\]]?$", text))


def add_body_paragraph(document: Document, chunks: list[str]) -> None:
    text = " ".join(chunks).strip()
    if text:
        document.add_paragraph(text, style="Lecture body")


def add_lecture_body(document: Document, fragments: list[str]) -> None:
    buffer: list[str] = []
    for text in fragments:
        buffer.append(text)
        if is_complete(text):
            add_body_paragraph(document, buffer)
            buffer.clear()
    add_body_paragraph(document, buffer)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source = Document(SOURCE)
    output = Document()
    section = output.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = output.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    body = output.styles.add_style("Lecture body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.15

    output.add_heading("Learning Prompt AI", level=0)
    output.add_paragraph("Cleaned lecture notes", style="Subtitle")
    output.add_paragraph(
        "Editorial cleanup: timing markers and transcript fragmentation have been removed. "
        "The original source file is unchanged; technical claims and examples remain source material to review.",
        style="Intense Quote",
    )

    events: list[tuple[str, int, list[str]]] = []
    current_part = 0
    current_lecture: int | None = None
    current_fragments: list[str] = []
    headings = 0
    original_nonempty = 0
    cleaned_fragments = 0

    for paragraph in source.paragraphs:
        raw = paragraph.text.strip()
        if not raw:
            continue
        original_nonempty += 1
        text = clean(raw)
        if not text:
            continue
        part = PART.match(text)
        if part:
            if current_lecture is not None:
                events.append((f"lecture:{current_part}:{current_lecture}", current_lecture, current_fragments))
            current_fragments = []
            current_part = int(part.group(1))
            events.append((f"part:{current_part}", current_part, []))
            current_lecture = None
            text = text[part.end():].strip(" -–—:")
        lecture = LECTURE.match(text)
        if lecture:
            if current_lecture is not None:
                events.append((f"lecture:{current_part}:{current_lecture}", current_lecture, current_fragments))
            current_lecture = int(lecture.group(1))
            current_fragments = []
            remainder = text[lecture.end():].strip(" -–—:")
            if remainder:
                current_fragments.append(remainder)
            continue
        if part:
            continue
        if current_lecture is None:
            continue
        current_fragments.append(text)
        cleaned_fragments += 1
    if current_lecture is not None:
        events.append((f"lecture:{current_part}:{current_lecture}", current_lecture, current_fragments))

    seen: set[str] = set()
    omitted_duplicates = 0
    for event, number, fragments in events:
        if event.startswith("part:"):
            output.add_heading(f"Part {number}", level=1)
            headings += 1
            continue
        key = SPACE.sub(" ", " ".join(fragments)).strip().lower()
        if len(key) > 1000 and key in seen:
            omitted_duplicates += 1
            continue
        seen.add(key)
        part_number = int(event.split(":")[1])
        title = TITLES.get((part_number, number), "")
        output.add_heading(f"Lecture {number}" + (f": {title}" if title else ""), level=2)
        headings += 1
        add_lecture_body(output, fragments)
    output.core_properties.title = "Learning Prompt AI - Lecture notes (cleaned)"
    output.core_properties.subject = "Cleaned transcript without time-related notes"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    output.save(OUTPUT)
    print(f"Created {OUTPUT} from {original_nonempty} source fragments; {headings} headings; {cleaned_fragments} cleaned fragments; {omitted_duplicates} duplicate lecture(s) omitted.")


if __name__ == "__main__":
    main()
